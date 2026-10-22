"""Builds the free sample MIS pack: assets/downloads/accuminds-monthly-mis-template.xlsx

    python tools/build_mis_template.py

All example figures are fictitious. Inputs are blue (and yellow where the user
should start); everything else is a formula, so the pack recalculates when the
inputs change. Recalculate once in Excel before publishing so previews show values.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

OUT = Path(__file__).resolve().parent.parent / "assets" / "downloads" / "accuminds-monthly-mis-template.xlsx"

FONT = "Arial"
NAVY, GOLD, SAND, GREY = "1A2E4A", "C9972B", "F7F5F1", "5B6778"
F = lambda **k: Font(name=FONT, **k)  # noqa: E731
INPUT = F(color="0000FF")
LINK = F(color="008000")
BOLD = F(bold=True)
HEAD = F(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor=SAND)
YELLOW = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="D9D4CA")
TOP = Border(top=Side(style="thin", color=NAVY))
NUM = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
DAYS = '0" days";(0" days");"-"'
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

# P&L lines: (label, kind) where kind is the section it adds to.
LINES = [
    ("Revenue", None),
    ("Product sales", "rev"), ("Service revenue", "rev"), ("Other operating income", "rev"),
    ("Cost of sales", None),
    ("Direct materials and purchases", "cos"), ("Direct labour", "cos"), ("Freight and other direct costs", "cos"),
    ("Operating expenses", None),
    ("Salaries and benefits", "opex"), ("Rent and utilities", "opex"), ("Sales and marketing", "opex"),
    ("Professional fees", "opex"), ("IT and software", "opex"), ("General and administrative", "opex"),
    ("Below EBITDA", None),
    ("Depreciation and amortisation", "dep"), ("Finance costs", "fin"), ("Corporate Tax (estimate)", "tax"),
]
BALANCE = ["Cash and bank", "Trade receivables", "Inventory", "Trade payables"]

# Fictitious monthly example values (AED): base per line, monthly growth/seasonality applied below.
BASE = {
    "Product sales": 420000, "Service revenue": 110000, "Other operating income": 6000,
    "Direct materials and purchases": 236000, "Direct labour": 38000, "Freight and other direct costs": 17000,
    "Salaries and benefits": 92000, "Rent and utilities": 21000, "Sales and marketing": 14000,
    "Professional fees": 7500, "IT and software": 4200, "General and administrative": 9800,
    "Depreciation and amortisation": 8500, "Finance costs": 3100, "Corporate Tax (estimate)": 6800,
}
SEASON = [0.92, 0.95, 1.04, 1.00, 1.06, 0.97, 0.88, 0.90, 1.03, 1.08, 1.10, 1.12]
BAL_BASE = {"Cash and bank": 610000, "Trade receivables": 455000, "Inventory": 300000, "Trade payables": 210000}


def example(label, m, budget):
    """Deterministic fictitious figure for a line and month (budget is smoother)."""
    base = BASE[label]
    fixed = label in ("Rent and utilities", "Depreciation and amortisation", "Finance costs", "IT and software")
    if budget:
        v = base * (1 if fixed else (0.97 + 0.01 * m))
    else:
        wobble = 1 + 0.035 * (((m * 7 + len(label)) % 5) - 2) / 2
        v = base * (1 if fixed else SEASON[m] * wobble)
    return round(v, -2)


def style_header(ws, row, first, last):
    for c in range(first, last + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill = HEAD, HEAD_FILL
        cell.alignment = Alignment(horizontal="center" if c > first else "left", vertical="center")


def setup(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = F(bold=True, size=16, color=NAVY)
    ws["A2"] = sub
    ws["A2"].font = F(size=10, color=GREY)


def data_sheet(wb, name, budget):
    """Actuals or Budget: line items down, Jan..Dec across (columns C..N)."""
    ws = wb.create_sheet(name)
    setup(ws, [34, 4] + [11] * 12 + [13])
    title(ws, f"{name} (AED)", "Blue = input. Replace the fictitious example figures with your own, month by month.")
    ws.cell(row=4, column=1, value="Line item")
    for i, m in enumerate(MONTHS):
        ws.cell(row=4, column=3 + i, value=m)
    ws.cell(row=4, column=15, value="Full year")
    style_header(ws, 4, 1, 15)
    rows = {}
    r = 5
    for label, kind in LINES:
        if kind is None:
            ws.cell(row=r, column=1, value=label).font = F(bold=True, color=NAVY)
            for c in range(1, 16):
                ws.cell(row=r, column=c).fill = SUB_FILL
        else:
            ws.cell(row=r, column=1, value=label).font = F()
            for m in range(12):
                cell = ws.cell(row=r, column=3 + m, value=example(label, m, budget))
                cell.font, cell.number_format = INPUT, NUM
            ws.cell(row=r, column=15, value=f"=SUM(C{r}:N{r})").number_format = NUM
            ws.cell(row=r, column=15).font = BOLD
            rows[label] = r
        r += 1
    if not budget:
        r += 1
        ws.cell(row=r, column=1, value="Balance sheet items (month-end)").font = F(bold=True, color=NAVY)
        for c in range(1, 16):
            ws.cell(row=r, column=c).fill = SUB_FILL
        r += 1
        for label in BALANCE:
            ws.cell(row=r, column=1, value=label).font = F()
            for m in range(12):
                v = round(BAL_BASE[label] * (0.94 + 0.012 * m) * (1 + 0.02 * ((m % 3) - 1)), -3)
                cell = ws.cell(row=r, column=3 + m, value=v)
                cell.font, cell.number_format = INPUT, NUM
            rows[label] = r
            r += 1
    ws.freeze_panes = "C5"
    return ws, rows


def build():
    wb = Workbook()

    # ---------------- Read me
    rm = wb.active
    rm.title = "Read me"
    setup(rm, [3, 26, 90])
    rm["B2"] = "Monthly MIS pack: sample template"
    rm["B2"].font = F(bold=True, size=18, color=NAVY)
    rm["B3"] = "By Accuminds Consulting FZCO · accumindsconsulting.ae · All figures are fictitious examples."
    rm["B3"].font = F(size=10, color=GREY)
    rows = [
        ("What it answers", "Are we on budget and why? Which lines drive profit? How healthy are cash and working capital? What needs a decision?"),
        ("1. Settings", "Enter your company name and pick the reporting month (1 = Jan … 12 = Dec)."),
        ("2. Actuals", "Enter actual results per month, plus month-end cash, receivables, inventory and payables."),
        ("3. Budget", "Enter your monthly budget for the same lines."),
        ("4. P&L vs Budget", "Calculated: the month and year-to-date against budget, with variances."),
        ("5. KPIs", "Calculated: margins, cost ratios and working-capital days for each month, with charts."),
        ("6. Cash outlook", "Enter expected receipts and payments for the next three months to see projected cash."),
        ("7. Commentary", "Write what changed, why, and the decisions needed. This is what turns numbers into a management pack."),
        ("Colour legend", "Blue text = input you can change. Yellow cell = start here. Black = formula, do not overwrite. Green = linked from another sheet."),
        ("Need help?", "We set up MIS packs and dashboards built from your own ledger: accumindsconsulting.ae/services/mis-reporting-dubai/"),
    ]
    for i, (k, v) in enumerate(rows, start=5):
        rm.cell(row=i, column=2, value=k).font = F(bold=True, color=NAVY)
        cell = rm.cell(row=i, column=3, value=v)
        cell.font, cell.alignment = F(), Alignment(wrap_text=True, vertical="top")
        rm.row_dimensions[i].height = 30

    # ---------------- Settings
    st = wb.create_sheet("Settings")
    setup(st, [3, 30, 32])
    title(st, "Settings", "Start here: yellow cells are the inputs this pack needs.")
    st["B4"], st["C4"] = "Company name", "Example Trading LLC (fictitious)"
    st["B5"], st["C5"] = "Reporting month (1-12)", 9
    st["B6"], st["C6"] = "Reporting month", '=INDEX({"Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"},C5)'
    st["B7"], st["C7"] = "Financial year", "2026"
    st["B8"], st["C8"] = "Days in a month (for KPI days)", 30
    st["B9"], st["C9"] = "Minimum cash buffer (AED)", 250000
    for r in (4, 5, 7, 8, 9):
        st.cell(row=r, column=3).font, st.cell(row=r, column=3).fill = INPUT, YELLOW
    for r in range(4, 10):
        st.cell(row=r, column=2).font = BOLD
    st["C9"].number_format = NUM
    st["C5"].comment = Comment("1 = January … 12 = December. P&L vs Budget and Cash outlook follow this month.", "Accuminds")
    st["C8"].comment = Comment("Assumption used for receivable, inventory and payable days. Change to 30.4 or actual days if you prefer.", "Accuminds")
    for nm, ref in (("RepMonth", "$C$5"), ("DaysInMonth", "$C$8"), ("MinCash", "$C$9")):
        wb.defined_names[nm] = DefinedName(nm, attr_text=f"Settings!{ref}")

    act, arow = data_sheet(wb, "Actuals", budget=False)
    bud, brow = data_sheet(wb, "Budget", budget=True)

    def month_val(sheet, r):
        return f"INDEX({sheet}!$C${r}:$N${r},RepMonth)"

    def ytd_val(sheet, r):
        return f"SUMPRODUCT((COLUMN({sheet}!$C${r}:$N${r})-COLUMN({sheet}!$C${r})+1<=RepMonth)*{sheet}!$C${r}:$N${r})"

    # ---------------- P&L vs Budget
    pl = wb.create_sheet("P&L vs Budget")
    setup(pl, [34, 13, 13, 13, 10, 3, 13, 13, 13, 10])
    pl["A1"] = '=Settings!C4&": P&L vs budget, "&Settings!C6&" "&Settings!C7'
    pl["A1"].font = F(bold=True, size=16, color=NAVY)
    pl["A2"] = "AED. Variance = actual minus budget; favourable is positive for income and profit, negative for costs."
    pl["A2"].font = F(size=10, color=GREY)
    hdr = ["", "Month actual", "Month budget", "Variance", "Var %", "", "YTD actual", "YTD budget", "Variance", "Var %"]
    for c, h in enumerate(hdr, start=1):
        pl.cell(row=4, column=c, value=h)
    style_header(pl, 4, 1, 10)

    out_rows = {}

    def put_line(r, label, a_expr, b_expr, ya_expr, yb_expr, bold=False, top=False):
        pl.cell(row=r, column=1, value=label).font = BOLD if bold else F()
        vals = {2: a_expr, 3: b_expr, 7: ya_expr, 8: yb_expr}
        for c, ex in vals.items():
            cell = pl.cell(row=r, column=c, value=f"={ex}")
            cell.number_format, cell.font = NUM, (BOLD if bold else LINK if ex.startswith(("INDEX", "SUMPRODUCT")) else F())
        for vc, ac, bc in ((4, "B", "C"), (9, "G", "H")):
            pl.cell(row=r, column=vc, value=f"={ac}{r}-{bc}{r}").number_format = NUM
            pl.cell(row=r, column=vc + 1, value=f'=IF({bc}{r}=0,"",{vc and get_column_letter(vc)}{r}/ABS({bc}{r}))').number_format = PCT
        if bold:
            for c in range(1, 11):
                pl.cell(row=r, column=c).font = BOLD
        if top:
            for c in range(1, 11):
                pl.cell(row=r, column=c).border = TOP
        out_rows[label] = r

    r = 5
    sections = {"rev": "Total revenue", "cos": "Total cost of sales", "opex": "Total operating expenses"}
    section_rows = {}
    current, first = None, None
    for label, kind in LINES:
        if kind is None:
            if current in sections:
                put_line(r, sections[current], f"SUM(B{first}:B{r-1})", f"SUM(C{first}:C{r-1})",
                         f"SUM(G{first}:G{r-1})", f"SUM(H{first}:H{r-1})", bold=True, top=True)
                section_rows[current] = r
                r += 1
                if current == "cos":
                    rv, cs = section_rows["rev"], section_rows["cos"]
                    put_line(r, "Gross profit", f"B{rv}-B{cs}", f"C{rv}-C{cs}", f"G{rv}-G{cs}", f"H{rv}-H{cs}", bold=True)
                    section_rows["gp"] = r
                    r += 1
                    pl.cell(row=r, column=1, value="Gross margin %").font = F(italic=True)
                    for c, col in ((2, "B"), (3, "C"), (7, "G"), (8, "H")):
                        cell = pl.cell(row=r, column=c, value=f'=IF({col}{rv}=0,"",{col}{section_rows["gp"]}/{col}{rv})')
                        cell.number_format, cell.font = PCT, F(italic=True)
                    r += 1
                if current == "opex":
                    gp, ox = section_rows["gp"], section_rows["opex"]
                    put_line(r, "EBITDA", f"B{gp}-B{ox}", f"C{gp}-C{ox}", f"G{gp}-G{ox}", f"H{gp}-H{ox}", bold=True)
                    section_rows["ebitda"] = r
                    r += 1
                r += 1
            pl.cell(row=r, column=1, value=label).font = F(bold=True, color=NAVY)
            for c in range(1, 11):
                pl.cell(row=r, column=c).fill = SUB_FILL
            r += 1
            current, first = None, r
            continue
        current = kind if kind in sections else current
        if kind in sections:
            current = kind
        put_line(r, label, month_val("Actuals", arow[label]), month_val("Budget", brow[label]),
                 ytd_val("Actuals", arow[label]), ytd_val("Budget", brow[label]))
        r += 1
    eb = section_rows["ebitda"]
    below = [out_rows["Depreciation and amortisation"], out_rows["Finance costs"], out_rows["Corporate Tax (estimate)"]]
    r += 0
    put_line(r, "Net profit", *[f"{c}{eb}-" + "-".join(f"{c}{b}" for b in below) for c in ("B", "C", "G", "H")], bold=True, top=True)
    np_row = r
    r += 1
    pl.cell(row=r, column=1, value="Net margin %").font = F(italic=True)
    for c, col in ((2, "B"), (3, "C"), (7, "G"), (8, "H")):
        cell = pl.cell(row=r, column=c, value=f'=IF({col}{section_rows["rev"]}=0,"",{col}{np_row}/{col}{section_rows["rev"]})')
        cell.number_format, cell.font = PCT, F(italic=True)
    # Cost lines: a positive variance is adverse, so colour by meaning.
    red = F(color="C0392B")
    green = F(color="1E8449")
    cost_rows = [out_rows[l] for l, k in LINES if k in ("cos", "opex", "dep", "fin", "tax")] + [section_rows["cos"], section_rows["opex"]]
    income_rows = [out_rows[l] for l, k in LINES if k == "rev"] + [section_rows["rev"], section_rows["gp"], section_rows["ebitda"], np_row]
    for rows_, good, bad in ((income_rows, ">", "<"), (cost_rows, "<", ">")):
        for rr in rows_:
            for col in ("D", "I"):
                pl.conditional_formatting.add(f"{col}{rr}", CellIsRule(operator="greaterThan" if good == ">" else "lessThan", formula=["0"], font=green))
                pl.conditional_formatting.add(f"{col}{rr}", CellIsRule(operator="lessThan" if good == ">" else "greaterThan", formula=["0"], font=red))
    pl.freeze_panes = "B5"

    # ---------------- KPIs (monthly trend from Actuals)
    kp = wb.create_sheet("KPIs")
    setup(kp, [34] + [11] * 12)
    kp["A1"] = '=Settings!C4&": KPIs by month"'
    kp["A1"].font = F(bold=True, size=16, color=NAVY)
    kp["A2"] = "Calculated from the Actuals sheet. Days use the 'Days in a month' setting."
    kp["A2"].font = F(size=10, color=GREY)
    kp.cell(row=4, column=1, value="KPI")
    for i, m in enumerate(MONTHS):
        kp.cell(row=4, column=2 + i, value=m)
    style_header(kp, 4, 1, 13)
    rev = [arow[l] for l, k in LINES if k == "rev"]
    cos = [arow[l] for l, k in LINES if k == "cos"]
    opx = [arow[l] for l, k in LINES if k == "opex"]

    def s(rows_, col):
        return "(" + "+".join(f"Actuals!{col}{x}" for x in rows_) + ")"

    kpis = [
        ("Revenue (AED)", lambda c: s(rev, c), NUM),
        ("Gross margin %", lambda c: f'IF({s(rev, c)}=0,"",1-{s(cos, c)}/{s(rev, c)})', PCT),
        ("Operating expenses % of revenue", lambda c: f'IF({s(rev, c)}=0,"",{s(opx, c)}/{s(rev, c)})', PCT),
        ("EBITDA margin %", lambda c: f'IF({s(rev, c)}=0,"",({s(rev, c)}-{s(cos, c)}-{s(opx, c)})/{s(rev, c)})', PCT),
        ("Receivable days (DSO)", lambda c: f'IF({s(rev, c)}=0,"",Actuals!{c}{arow["Trade receivables"]}/{s(rev, c)}*DaysInMonth)', DAYS),
        ("Inventory days", lambda c: f'IF({s(cos, c)}=0,"",Actuals!{c}{arow["Inventory"]}/{s(cos, c)}*DaysInMonth)', DAYS),
        ("Payable days (DPO)", lambda c: f'IF({s(cos, c)}=0,"",Actuals!{c}{arow["Trade payables"]}/{s(cos, c)}*DaysInMonth)', DAYS),
        ("Cash conversion cycle", None, DAYS),
        ("Month-end cash (AED)", lambda c: f'Actuals!{c}{arow["Cash and bank"]}', NUM),
    ]
    for i, (label, fn, fmt) in enumerate(kpis, start=5):
        kp.cell(row=i, column=1, value=label).font = BOLD
        for m in range(12):
            col_src = get_column_letter(3 + m)
            col_dst = get_column_letter(2 + m)
            if fn is None:
                f = f'=IF(OR({col_dst}9="",{col_dst}10="",{col_dst}11=""),"",{col_dst}9+{col_dst}10-{col_dst}11)'
            else:
                f = "=" + fn(col_src)
            cell = kp.cell(row=i, column=2 + m, value=f)
            cell.number_format = fmt
    kp.freeze_panes = "B5"
    ch = BarChart()
    ch.title, ch.y_axis.title, ch.height, ch.width = "Revenue by month (AED)", None, 7, 18
    ch.add_data(Reference(kp, min_col=1, max_col=13, min_row=5, max_row=5), from_rows=True, titles_from_data=True)
    ch.set_categories(Reference(kp, min_col=2, max_col=13, min_row=4, max_row=4))
    kp.add_chart(ch, "A16")
    lc = LineChart()
    lc.title, lc.height, lc.width = "Gross and EBITDA margin", 7, 18
    lc.add_data(Reference(kp, min_col=1, max_col=13, min_row=6, max_row=6), from_rows=True, titles_from_data=True)
    lc.add_data(Reference(kp, min_col=1, max_col=13, min_row=8, max_row=8), from_rows=True, titles_from_data=True)
    lc.set_categories(Reference(kp, min_col=2, max_col=13, min_row=4, max_row=4))
    lc.y_axis.number_format = "0%"
    kp.add_chart(lc, "H16")

    # ---------------- Cash outlook (next three months)
    co = wb.create_sheet("Cash outlook")
    setup(co, [36, 15, 15, 15])
    co["A1"] = '=Settings!C4&": cash outlook, next three months"'
    co["A1"].font = F(bold=True, size=16, color=NAVY)
    co["A2"] = "AED. Enter expected receipts and payments (blue). Opening cash links to month-end cash in Actuals."
    co["A2"].font = F(size=10, color=GREY)
    for c, h in enumerate(["", "Month +1", "Month +2", "Month +3"], start=1):
        co.cell(row=4, column=c, value=h)
    style_header(co, 4, 1, 4)
    co["A5"] = "Opening cash"
    co["B5"] = f"=INDEX(Actuals!$C${arow['Cash and bank']}:$N${arow['Cash and bank']},RepMonth)"
    co["B5"].font = LINK
    co["C5"], co["D5"] = "=B16", "=C16"
    receipts = [("Receipts from customers", 548000), ("Other receipts", 5000)]
    payments = [("Payments to suppliers", 300000), ("Salaries and benefits", 94000), ("Rent and utilities", 21000),
                ("Other operating costs", 37000), ("VAT payable to FTA", 22000), ("Corporate Tax", 0),
                ("Loan repayments and finance costs", 15000)]
    r = 6
    for label, v in receipts:
        co.cell(row=r, column=1, value=label)
        for c in (2, 3, 4):
            cell = co.cell(row=r, column=c, value=round(v * (1 + 0.02 * (c - 2)), -3))
            cell.font, cell.number_format = INPUT, NUM
        r += 1
    rec_end = r - 1
    for label, v in payments:
        co.cell(row=r, column=1, value=label)
        for c in (2, 3, 4):
            val = round(v * (1 + 0.015 * (c - 2)), -3)
            if label == "Corporate Tax" and c == 4:
                val = 20000
            cell = co.cell(row=r, column=c, value=val)
            cell.font, cell.number_format = INPUT, NUM
        r += 1
    pay_end = r - 1
    co.cell(row=r, column=1, value="Net cash flow").font = BOLD
    for c in (2, 3, 4):
        L = get_column_letter(c)
        co.cell(row=r, column=c, value=f"=SUM({L}6:{L}{rec_end})-SUM({L}{rec_end+1}:{L}{pay_end})").number_format = NUM
        co.cell(row=r, column=c).font = BOLD
    net = r
    r += 1
    assert r == 16, r
    co.cell(row=r, column=1, value="Closing cash").font = BOLD
    for c in (2, 3, 4):
        L = get_column_letter(c)
        cell = co.cell(row=r, column=c, value=f"={L}5+{L}{net}")
        cell.number_format, cell.font, cell.border = NUM, BOLD, TOP
    r += 1
    co.cell(row=r, column=1, value="Above minimum cash buffer?")
    for c in (2, 3, 4):
        L = get_column_letter(c)
        co.cell(row=r, column=c, value=f'=IF({L}16>=MinCash,"Yes","Below buffer")')
    for c in range(2, 5):
        co.cell(row=5, column=c).number_format = NUM
    co.conditional_formatting.add("B17:D17", CellIsRule(operator="equal", formula=['"Below buffer"'], font=F(bold=True, color="C0392B")))

    # ---------------- Commentary
    cm = wb.create_sheet("Commentary")
    setup(cm, [3, 30, 90])
    cm["B2"] = '=Settings!C4&": management commentary, "&Settings!C6&" "&Settings!C7'
    cm["B2"].font = F(bold=True, size=16, color=NAVY)
    prompts = [
        ("Headline", "One or two sentences: how did the month go against budget?"),
        ("Revenue", "What drove the variance: volume, price, mix, timing?"),
        ("Margins and costs", "Which costs moved and why? Any one-offs?"),
        ("Cash and working capital", "Receivable, inventory and payable days: any customers or stock to act on?"),
        ("Risks and opportunities", "What could change the next three months?"),
        ("Decisions needed", "What do owners or managers need to decide this month?"),
    ]
    for i, (k, v) in enumerate(prompts, start=4):
        cm.cell(row=i, column=2, value=k).font = F(bold=True, color=NAVY)
        cell = cm.cell(row=i, column=3, value=v)
        cell.font, cell.fill = F(color="0000FF", italic=True), YELLOW
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cm.row_dimensions[i].height = 48

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = GOLD if ws.title in ("Settings", "Read me") else NAVY
        for row in ws.iter_rows():
            for cell in row:
                if cell.font and cell.font.name != FONT:
                    cell.font = Font(name=FONT, bold=cell.font.bold, italic=cell.font.italic,
                                     color=cell.font.color, size=cell.font.size)
    wb.calculation.fullCalcOnLoad = True
    wb.properties.creator = "Accuminds Consulting FZCO"
    wb.properties.title = "Monthly MIS pack: sample template"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
